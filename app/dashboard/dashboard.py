from __future__ import annotations

import io
import json
import os
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pandas as pd
import requests
import streamlit as st
from docx import Document

try:
    from app.services.llm.provider import LOG_PATH, LLMService
except Exception:
    LLMService = None
    LOG_PATH = Path(
        os.getenv(
            "LLM_PROCESSING_LOG",
            str(Path.home() / ".resume_hub_llm.log.jsonl"),
        )
    )

# ==============================================================================
# CONFIGURATION
# ==============================================================================

_configured_api_base = os.getenv("FASTAPI_API_BASE", "").strip()
DEFAULT_API_BASE = (
    _configured_api_base.rstrip("/") if _configured_api_base else "http://localhost:8000"
)
REQUEST_TIMEOUT = 120

PROVIDER_LABELS = {
    "gemini": "Google Gemini",
    "groq": "Groq",
    "openrouter": "OpenRouter",
    "deepseek": "DeepSeek",
    "openai": "OpenAI",
    "claude": "Anthropic Claude",
    "ollama": "Ollama",
    "experiential": "Experiential Cloud",
}

DIRECT_DEFAULT_MODELS = {
    "gemini": "gemini-3.7-flash",
    "groq": "qwen3.8-27b",
    "openrouter": "deepseek-v4-flash",
    "deepseek": "deepseek-v4-flash",
    "openai": "gpt-5.6-luna",
    "claude": "claude-fable-5",
    "ollama": "qwen3.8-27b",
}

DIRECT_PROVIDER_MODELS = {
    "gemini": [
        "gemini-3.7-flash",
        "gemini-3.6-flash",
        "gemini-3.5-flash",
        "gemini-3.1-pro-preview",
        "gemini-3.1-flash-lite",
        "gemini-2.5-flash",
        "gemini-2.5-pro",
        "gemini-2.5-flash-lite",
    ],
    "groq": ["qwen3.8-27b", "llama-3.3-70b-versatile"],
    "openrouter": ["deepseek-v4-flash", "deepseek-r1", "gemini-3.7-flash"],
    "deepseek": [
        "deepseek-v4-flash",
        "deepseek-chat",
        "deepseek-coder",
        "deepseek-r1",
    ],
    "openai": ["gpt-5.6-luna", "gpt-4o", "gpt-4o-mini", "o3-mini"],
    "claude": [
        "claude-fable-5",
        "claude-3-5-sonnet-20241022",
        "claude-3-5-haiku-20241022",
    ],
    "ollama": ["qwen3.8-27b", "llama3.2", "mistral", "deepseek-r1"],
}


def direct_model_options(provider: str) -> list[str]:
    defaults = DIRECT_PROVIDER_MODELS.get(provider, [DIRECT_DEFAULT_MODELS.get(provider, "")])
    env_name = f"DIRECT_MODELS_{provider.upper()}"
    configured = [x.strip() for x in os.getenv(env_name, "").split(",") if x.strip()]

    options = []
    for model in [*defaults, *configured]:
        if model and model not in options:
            options.append(model)
    return options or [DIRECT_DEFAULT_MODELS.get(provider, "")]


TEMPLATE_LABELS = {
    "auto": "🎯 Auto-detect (match JD language)",
    "international_ats": "International English ATS (Single-Page)",
    "academic": "Academic / Research (Serif, Education-first)",
    "technical_lead": "Technical Lead (Open Source + Speaking)",
    "hr_executive_gold": "HR Gold Standard (Executive)",
    "german_corporate": "Corporate Slate Navy",
    "german_minimal_ats": "German Minimal ATS (Single-Column)",
    "german_modern": "Modern Two-Column",
    "german_classic": "German Classic Single-Column PDF",
    "standard": "Standard ATS Single-Column",
}

st.set_page_config(
    page_title="AI Resume & CV Optimization Hub",
    page_icon="🎯",
    layout="wide",
)

# ==============================================================================
# API ENDPOINTS
# ==============================================================================


def endpoints(api_base: str) -> dict[str, str]:
    prefix = f"{api_base.rstrip('/')}/api/v1/resume"
    return {
        "analyze": f"{prefix}/analyze",
        "full_docx": f"{prefix}/generate-full",
        "german_pdf": f"{prefix}/generate-german-cv",
        "german_tex": f"{prefix}/generate-tex-cv",
        "diff_preview": f"{prefix}/diff-preview",
        "analyze_bulk": f"{prefix}/analyze-bulk",
        "audit_matrix": f"{prefix}/audit-matrix",
        "cover_letter": f"{prefix}/generate-cover-letter",
        "interview_prep": f"{prefix}/interview-prep",
        "linkedin_optimize": f"{prefix}/linkedin-optimize",
        "tracker": f"{prefix}/tracker/applications",
        "usage_summary": f"{prefix}/usage-summary",
        "model_catalog": f"{prefix}/model-catalog",
        "quota_status": f"{prefix}/quota-status",
        "quota_events": f"{prefix}/quota-events",
        "processing_log": f"{prefix}/processing-log",
    }


# ==============================================================================
# SESSION STATE
# ==============================================================================

st.session_state.setdefault("api_base", DEFAULT_API_BASE)
st.session_state.setdefault("route_mode", "experiential")
st.session_state.setdefault("experiential_model", "gpt-6-astra")
st.session_state.setdefault("direct_provider", "gemini")
st.session_state.setdefault("direct_model", DIRECT_DEFAULT_MODELS["gemini"])
st.session_state.setdefault("last_analysis", None)
st.session_state.setdefault("uploaded_file_data", None)
st.session_state.setdefault("job_desc", "")

# ==============================================================================
# STYLING
# ==============================================================================

st.markdown(
    """
<style>
.block-container { padding-top: 2rem; padding-bottom: 4rem; max-width: 1450px; }
h1 { letter-spacing: -0.035em; } h2, h3 { letter-spacing: -0.02em; }
[data-testid="stSidebar"] { border-right: 1px solid rgba(128,128,128,.18); }
[data-testid="stSidebar"] .block-container { padding-top: 1.5rem; }
.ui-card { border:1px solid rgba(128,128,128,.18); border-radius:14px; padding:14px 16px; background:rgba(128,128,128,.035); min-height:86px; }
.ui-card-title { font-size:.74rem; text-transform:uppercase; letter-spacing:.07em; color:#6b7280; font-weight:700; margin-bottom:5px; }
.ui-card-value { font-size:1.2rem; font-weight:750; line-height:1.15; }
.ui-card-sub { font-size:.8rem; color:#6b7280; margin-top:5px; }
.metric-big { font-size:2.25rem; font-weight:750; line-height:1; }
.metric-label { color:#6B7280; font-size:.78rem; text-transform:uppercase; letter-spacing:.06em; }
.workflow { display:flex; align-items:center; gap:8px; margin:0 0 22px; color:#6b7280; font-size:.86rem; }
.workflow-step { display:flex; align-items:center; gap:7px; white-space:nowrap; font-weight:650; }
.step-badge { display:inline-flex; align-items:center; justify-content:center; width:26px; height:26px; border-radius:50%; background:#e5e7eb; color:#374151; font-weight:700; font-size:.78rem; }
.step-badge.active { background:#2563eb; color:white; } .step-badge.done { background:#16a34a; color:white; }
.workflow-line { height:1px; flex:0 1 70px; background:rgba(128,128,128,.3); }
.chip { display:inline-block; padding:4px 10px; margin:2px 4px 2px 0; border-radius:999px; font-size:.8rem; font-weight:600; }
.chip-good { background:#dcfce7; color:#166534; } .chip-bad { background:#fee2e2; color:#991b1b; }
.model-free { color:#15803d; font-weight:700; } .model-promo { color:#2563eb; font-weight:700; } .model-paid { color:#6b7280; font-weight:600; }
.provider-card { padding:10px 12px; border-radius:10px; margin-bottom:7px; border:1px solid rgba(128,128,128,.18); }
.security-note { padding:13px; border-radius:10px; background:#1e293b; border:1px solid #334155; color:#f8fafc; font-size:.86rem; }
.quota-card { border:1px solid rgba(128,128,128,.18); border-radius:12px; padding:12px 14px; margin:7px 0; }
.quota-title { font-weight:700; } .quota-meta { color:#6b7280; font-size:.78rem; }
.quota-bar-bg { background:#1e293b; border-radius:999px; height:7px; overflow:hidden; }
.quota-bar-fill { height:100%; border-radius:999px; }
.health-ok { color:#15803d; font-weight:700; } .health-warn { color:#b45309; font-weight:700; } .health-bad { color:#b91c1c; font-weight:700; }
button[kind="primary"] { min-height:44px; }
[data-testid="stDataFrame"] { border-radius:10px; overflow:hidden; }
</style>
""",
    unsafe_allow_html=True,
)

# ==============================================================================
# HELPERS
# ==============================================================================


def get_experiential_catalog() -> list[dict]:
    if LLMService and hasattr(LLMService, "get_model_catalog"):
        try:
            rows = LLMService.get_model_catalog("experiential")
            if rows:
                return rows
        except Exception:
            pass
    return [
        {
            "model": "gpt-6-astra",
            "label": "GPT-6 Astra",
            "provider": "experiential",
            "status": "PROMOTION",
            "free": True,
            "promotion": True,
        },
        {
            "model": "gpt-5.6-luna",
            "label": "GPT-5.6 Luna",
            "provider": "experiential",
            "status": "PROMOTION",
            "free": True,
            "promotion": True,
        },
        {
            "model": "deepseek-v4-flash",
            "label": "DeepSeek V4 Flash",
            "provider": "experiential",
            "status": "PROMOTION",
            "free": True,
            "promotion": True,
        },
        {
            "model": "qwen3.8-27b",
            "label": "Qwen3.8 27B",
            "provider": "experiential",
            "status": "PROMOTION",
            "free": True,
            "promotion": True,
        },
        {
            "model": "gemini-3.7-flash",
            "label": "Gemini 3.7 Flash",
            "provider": "experiential",
            "status": "FREE",
            "free": True,
            "promotion": False,
        },
    ]


def model_display(row: dict) -> str:
    label = row.get("label") or row.get("model", "")
    status = str(row.get("status", "PAID")).upper()
    if status == "PROMOTION":
        return f"{label}  ·  FREE PROMOTION"
    if status == "FREE":
        return f"{label}  ·  FREE"
    return f"{label}  ·  PAID"


def current_llm_config() -> dict:
    mode = st.session_state.get("route_mode", "experiential")
    if mode == "experiential":
        model = st.session_state.get("experiential_model", "gpt-6-astra")
        row = next(
            (r for r in get_experiential_catalog() if r.get("model") == model),
            {},
        )
        return {
            "route_mode": "experiential",
            "provider": "experiential",
            "model_name": model,
            "status": row.get("status", "UNKNOWN"),
            "label": row.get("label", model),
        }
    provider = st.session_state.get("direct_provider", "gemini")
    model_name = st.session_state.get("direct_model") or DIRECT_DEFAULT_MODELS.get(provider, "")
    return {
        "route_mode": "direct",
        "provider": provider,
        "model_name": model_name,
        "status": "DIRECT API",
        "label": f"{PROVIDER_LABELS.get(provider, provider)} — {model_name}",
    }


_ACTIONABLE_PREFIXES = (
    "CRITICAL MATCH GAP:",
    "MODERATE MATCH GAP:",
    "GOOD ALIGNMENT:",
    "HIGH ALIGNMENT:",
    "KEYWORD INJECTION:",
    "FULL KEYWORD COVERAGE:",
    "LANGUAGE MISMATCH",
)

_META_PREFIXES = (
    "VERIFICATION:",
    "FORMATTING RULE",
    "OUTSTANDING ATS MATCH:",
)


def actionable_suggestions_from_last_analysis() -> list[str]:
    result = st.session_state.get("last_analysis") or {}
    raw = result.get("improvement_suggestions") or []
    out: list[str] = []
    for rec in raw:
        text = (rec or "").strip()
        if not text:
            continue
        if any(text.startswith(p) for p in _META_PREFIXES):
            continue
        out.append(text)
    return out


def llm_request_data(**extra) -> dict:
    data = dict(extra)
    data.update(current_llm_config())

    suggestions = actionable_suggestions_from_last_analysis()
    if suggestions:
        data["improvement_suggestions"] = json.dumps(suggestions)

    return data


# ---------------------------------------------------------------------------
# Usage, catalog, quota helpers
# ---------------------------------------------------------------------------


@st.cache_data(ttl=15)
def fetch_usage_summary(api_base: str, period: str = "all") -> dict | None:
    try:
        url = f"{api_base.rstrip('/')}/api/v1/resume/usage-summary"
        r = requests.get(url, params={"period": period}, timeout=10)
        if r.status_code == 200:
            return r.json().get("summary")
    except Exception:
        return None
    return None


@st.cache_data(ttl=15)
def fetch_model_catalog(api_base: str, provider: str = "experiential") -> list[dict]:
    try:
        url = f"{api_base.rstrip('/')}/api/v1/resume/model-catalog"
        r = requests.get(url, params={"provider": provider}, timeout=10)
        if r.status_code == 200:
            return r.json().get("models") or []
    except Exception:
        pass

    if LLMService and hasattr(LLMService, "get_model_catalog"):
        try:
            return LLMService.get_model_catalog(provider)
        except Exception:
            pass
    return []


@st.cache_data(ttl=300)
def fetch_career_options(api_base: str) -> dict:
    """Static catalog of cover-letter templates and interview families."""
    try:
        url = f"{api_base.rstrip('/')}/api/v1/resume/career-options"
        r = requests.get(url, timeout=10)
        if r.status_code == 200:
            return r.json()
    except Exception:
        pass
    return {"cover_letter_templates": [], "interview_families": []}


@st.cache_data(ttl=15)
def fetch_quota_status(api_base: str) -> dict | None:
    try:
        url = f"{api_base.rstrip('/')}/api/v1/resume/quota-status"
        r = requests.get(url, timeout=10)
        if r.status_code == 200:
            return r.json().get("providers") or {}
    except Exception:
        return None
    return None


@st.cache_data(ttl=5)
def fetch_processing_log(api_base: str, limit: int = 50) -> list[dict]:
    """Fetch recent backend processing-log events."""
    try:
        url = f"{api_base.rstrip('/')}/api/v1/resume/processing-log"
        r = requests.get(url, params={"limit": limit}, timeout=10)
        if r.status_code == 200:
            return r.json().get("logs") or []
    except Exception:
        pass
    return []


def _fmt_log_detail(e: dict) -> str:
    """Compact one-line summary of a pipeline log event."""
    parts: list[str] = []
    if e.get("provider"):
        mdl = e.get("model") or ""
        parts.append(f"{e['provider']}/{mdl}" if mdl else str(e["provider"]))
    if e.get("operation"):
        parts.append(f"op={e['operation']}")
    if e.get("layout"):
        parts.append(f"layout={e['layout']}")
    if e.get("file_type"):
        parts.append(f"file={e['file_type']}")
    if e.get("chars_extracted") is not None:
        parts.append(f"chars={e['chars_extracted']}")
    if e.get("ats_score") is not None:
        parts.append(f"ats={e['ats_score']}")
    if e.get("missing_count") is not None:
        parts.append(f"missing={e['missing_count']}")
    if e.get("total_tokens"):
        parts.append(f"tok={e['total_tokens']}")
    if e.get("pages"):
        parts.append(f"pages={e['pages']}")
    if e.get("bytes"):
        parts.append(f"{e['bytes']}B")
    if e.get("cache"):
        parts.append(f"cache={e['cache']}")
    return " · ".join(parts)


def status_badge(status: str) -> str:
    s = (status or "").upper()
    if s == "PROMOTION":
        return '<span class="model-promo">FREE PROMOTION</span>'
    if s == "FREE":
        return '<span class="model-free">FREE</span>'
    return '<span class="model-paid">PAID</span>'


def fmt_tokens(n: int) -> str:
    try:
        n = int(n or 0)
    except Exception:
        return "0"
    if n >= 1_000_000:
        return f"{n / 1_000_000:.2f}M"
    if n >= 1_000:
        return f"{n / 1_000:.1f}K"
    return str(n)


def fmt_reset_local(iso: str) -> str:
    if not iso:
        return "—"
    try:
        dt = datetime.fromisoformat(iso.replace("Z", "+00:00"))
    except Exception:
        return iso
    now = datetime.now(timezone.utc)
    if (dt - now) < timedelta(days=1):
        return dt.astimezone().strftime("%H:%M")
    return dt.astimezone().strftime("%d %b at %H:%M")


def quota_bar(percent: float) -> str:
    pct = max(0.0, min(100.0, float(percent or 0.0)))
    color = "#10b981" if pct < 60 else "#f59e0b" if pct < 85 else "#ef4444"
    return (
        '<div style="background:#1e293b;border-radius:6px;height:8px;'
        'overflow:hidden;margin:6px 0 4px 0;">'
        f'<div style="width:{pct:.1f}%;height:100%;background:{color};'
        'transition:width .3s ease;"></div>'
        "</div>"
    )


def chips(items, kind="good"):
    if not items:
        st.caption("None found.")
        return
    css = "chip-good" if kind == "good" else "chip-bad"
    st.markdown(
        "".join(f'<span class="chip {css}">{item}</span>' for item in items),
        unsafe_allow_html=True,
    )


def score_block(label: str, value: int):
    value = max(0, min(int(value or 0), 100))
    st.markdown(f'<div class="metric-label">{label}</div>', unsafe_allow_html=True)
    st.markdown(f'<div class="metric-big">{value}%</div>', unsafe_allow_html=True)
    st.progress(value / 100)


def error_detail(response) -> str:
    if response is None:
        return "Could not connect to the backend service."
    try:
        data = response.json()
        if isinstance(data, dict) and data.get("detail"):
            return str(data["detail"])
        return response.text
    except Exception:
        return response.text or f"HTTP {response.status_code}"


def api_request(url: str, data: dict, files, spinner_text: str):
    with st.status(spinner_text, expanded=True) as status:
        st.write("Sending payload and document to backend...")
        try:
            response = requests.post(url, data=data, files=files, timeout=REQUEST_TIMEOUT)
        except requests.exceptions.ConnectionError:
            status.update(label="Backend unavailable.", state="error")
            return None
        except requests.exceptions.Timeout:
            status.update(label="Backend request timed out.", state="error")
            return None
        except requests.exceptions.RequestException as exc:
            status.update(label="Request failed.", state="error")
            st.error(str(exc))
            return None

        if response.status_code in (429, 503):
            wait_time = 10 if response.status_code == 429 else 5
            status.update(
                label=f"Server busy ({response.status_code}). Retrying in {wait_time}s...",
                state="running",
            )
            progress = st.progress(0)
            for i in range(wait_time):
                time.sleep(1)
                progress.progress((i + 1) / wait_time)
            progress.empty()
            try:
                response = requests.post(url, data=data, files=files, timeout=REQUEST_TIMEOUT)
            except requests.exceptions.RequestException:
                status.update(label="Retry failed.", state="error")
                return None

        if response.status_code == 200:
            status.update(label="Analysis complete.", state="complete", expanded=False)
        else:
            status.update(label=f"Request failed ({response.status_code})", state="error")
    return response


def build_files_payload():
    uploaded = st.session_state.get("uploaded_file_data")
    if not uploaded:
        return {}
    filename, file_bytes, mime_type = uploaded
    return {"resume_file": (filename, file_bytes, mime_type)}


def reset_all():
    st.session_state["last_analysis"] = None
    st.session_state["uploaded_file_data"] = None
    st.session_state["job_desc"] = ""


def generate_txt_export(result_data: dict) -> bytes:
    lines = [
        "=" * 60,
        "AI RESUME & CV OPTIMIZATION REPORT".center(60),
        "=" * 60,
        "",
        f"ATS Match Score: {result_data.get('ats_match_score', 0)}%",
        f"Keyword Density Score: {result_data.get('keyword_density_score', 0)}%",
        "",
        "[MATCHING SKILLS]",
        ", ".join(result_data.get("matching_skills", [])) or "None",
        "",
        "[MISSING SKILLS]",
        ", ".join(result_data.get("missing_skills", [])) or "None",
        "",
        "=" * 60,
    ]
    for index, suggestion in enumerate(result_data.get("improvement_suggestions", []), 1):
        lines.extend([f"{index}. {suggestion}", ""])
    return "\n".join(lines).encode("utf-8")


def generate_docx_export(result_data: dict) -> bytes:
    doc = Document()
    doc.add_heading("AI Resume Optimization Report", level=0)
    doc.add_heading("Overview Metrics", level=1)
    p = doc.add_paragraph()
    p.add_run("ATS Match Score: ").bold = True
    p.add_run(f"{result_data.get('ats_match_score', 0)}%\n")
    p.add_run("Keyword Density: ").bold = True
    p.add_run(f"{result_data.get('keyword_density_score', 0)}%")
    doc.add_heading("Skill Gap Analysis", level=1)
    p = doc.add_paragraph()
    p.add_run("Matching Skills: ").bold = True
    p.add_run(", ".join(result_data.get("matching_skills", [])) or "None")
    p.add_run("\n")
    p.add_run("Missing Skills: ").bold = True
    p.add_run(", ".join(result_data.get("missing_skills", [])) or "None")
    doc.add_heading("Suggestions & Rewritten Bullets", level=1)
    for suggestion in result_data.get("improvement_suggestions", []):
        doc.add_paragraph(suggestion, style="List Bullet")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def render_diff_view(diff_data_list):
    st.markdown("### Visual Bullet Point Comparison")
    for index, item in enumerate(diff_data_list, 1):
        similarity = item.get("similarity_ratio", 0)
        with st.expander(
            f"Bullet Point #{index} (Similarity: {similarity}%)",
            expanded=True,
        ):
            col1, col2 = st.columns(2)
            with col1:
                st.caption("Original Bullet Point")
                st.info(item.get("original") or "N/A")
            with col2:
                st.caption("Optimized Bullet Point")
                st.success(item.get("optimized") or "N/A")
            st.caption("Inline Changes:")
            st.markdown(
                item.get("diff_html", ""),
                unsafe_allow_html=True,
            )


# ==============================================================================
# SIDEBAR — ROUTING (fragment)
# ==============================================================================


@st.fragment
def _render_sidebar():
    with st.sidebar:
        st.title("AI Engine")

        route_options = ["experiential", "direct"]
        route_labels = {
            "experiential": "Experiential Cloud",
            "direct": "Direct API",
        }
        selected_route = st.selectbox(
            "Routing",
            route_options,
            index=route_options.index(st.session_state["route_mode"]),
            format_func=lambda x: route_labels[x],
        )
        st.session_state["route_mode"] = selected_route

        if selected_route == "experiential":
            catalog = get_experiential_catalog()
            slugs = [row["model"] for row in catalog]
            current = st.session_state.get("experiential_model", "gpt-6-astra")
            if current not in slugs:
                current = slugs[0] if slugs else "gpt-6-astra"
            selected_model = st.selectbox(
                "Experiential Model",
                slugs,
                index=slugs.index(current) if current in slugs else 0,
                format_func=lambda slug: model_display(
                    next(r for r in catalog if r["model"] == slug)
                ),
            )
            st.session_state["experiential_model"] = selected_model
            selected_row = next((r for r in catalog if r["model"] == selected_model), {})
            status = selected_row.get("status", "UNKNOWN")
            if status == "PROMOTION":
                st.markdown(
                    '<span class="model-promo">FREE PROMOTION</span>',
                    unsafe_allow_html=True,
                )
            elif status == "FREE":
                st.markdown('<span class="model-free">FREE</span>', unsafe_allow_html=True)
            else:
                st.markdown('<span class="model-paid">PAID</span>', unsafe_allow_html=True)
            st.caption("Lane: Experiential Cloud")
        else:
            provider_options = [p for p in PROVIDER_LABELS if p != "experiential"]
            current_provider = st.session_state.get("direct_provider", "gemini")
            provider = st.selectbox(
                "Direct API Provider",
                provider_options,
                index=(
                    provider_options.index(current_provider)
                    if current_provider in provider_options
                    else 0
                ),
                format_func=lambda x: PROVIDER_LABELS[x],
            )
            st.session_state["direct_provider"] = provider

            options = direct_model_options(provider)
            current_model = st.session_state.get("direct_model") or options[0]
            if current_model not in options:
                options = [current_model, *options]

            direct_model = st.selectbox(
                "Direct API Model",
                options,
                index=(options.index(current_model) if current_model in options else 0),
                help="Provider-specific model list. Add extra models with DIRECT_MODELS_<PROVIDER> in the dashboard environment.",
            )
            st.session_state["direct_model"] = direct_model
            st.caption(
                "Lane: Direct Provider API. Calls bypass gateway and hit native endpoints directly."
            )

        cfg = current_llm_config()
        st.divider()
        st.markdown(f"**Active route:** `{cfg['route_mode']}`")
        st.markdown(f"**Provider:** `{cfg['provider']}`")
        st.markdown(f"**Model:** `{cfg['model_name']}`")
        st.markdown(f"**Status:** `{cfg['status']}`")

        st.markdown(
            """
            <div class="security-note">
            <strong>API keys are backend-only.</strong><br><br>
            The dashboard does not request, store, or transmit provider keys.
            Credentials remain in the backend environment.
            </div>
            """,
            unsafe_allow_html=True,
        )

        with st.expander("Backend configuration", expanded=False):
            api_base_input = st.text_input(
                "FastAPI Backend URL",
                value=st.session_state.get("api_base", DEFAULT_API_BASE),
            )
            if api_base_input.strip():
                st.session_state["api_base"] = api_base_input.strip().rstrip("/")
            st.caption("API keys are configured in the backend environment.")

            if LLMService:
                try:
                    for item in LLMService.provider_status():
                        configured = item.get("configured", False)
                        symbol = "✓" if configured else "○"
                        name = item.get("provider", "")
                        st.markdown(
                            f'<div class="provider-card"><strong>{symbol} {PROVIDER_LABELS.get(name, name)}</strong><br><small>Model: {item.get("model", "")}<br>Auth: {item.get("authentication", "")}</small></div>',
                            unsafe_allow_html=True,
                        )
                except Exception as exc:
                    st.warning(f"Could not read provider status: {exc}")

            st.markdown("---")
            st.markdown("**Experiential Model Catalog**")
            catalog = fetch_model_catalog(
                st.session_state.get("api_base", DEFAULT_API_BASE),
                "experiential",
            )
            if not catalog:
                st.caption("Could not load model catalog.")
            else:
                for row in catalog:
                    model_id = row.get("model", "")
                    label = row.get("label") or model_id
                    status = row.get("status", "PAID")
                    badge = status_badge(status)
                    in_per_k = row.get("input_per_1k", 0.0)
                    out_per_k = row.get("output_per_1k", 0.0)
                    if row.get("free"):
                        pricing_line = "Free — no billing"
                    else:
                        pricing_line = f"${in_per_k:.4f} in / ${out_per_k:.4f} out per 1K"
                    st.markdown(
                        f'<div class="provider-card">'
                        f"<strong>{label}</strong> &nbsp; {badge}<br>"
                        f"<small><code>{model_id}</code><br>"
                        f"{pricing_line}</small></div>",
                        unsafe_allow_html=True,
                    )

        st.caption(f"Backend: {st.session_state.get('api_base')}")
        if st.session_state.get("last_analysis") and st.button("Start Over", width="stretch"):
            reset_all()
            st.rerun()


_render_sidebar()

# ==============================================================================
# ACTIVE API CONFIG
# ==============================================================================

ACTIVE_API_BASE = st.session_state["api_base"].strip().rstrip("/")
ENDPOINTS = endpoints(ACTIVE_API_BASE)
cfg = current_llm_config()

# ==============================================================================
# HEADER
# ==============================================================================

has_analysis = bool(st.session_state.get("last_analysis"))
step = 2 if has_analysis else 1

st.title("AI Resume & CV Matcher")
st.caption(
    "One workspace for ATS analysis, targeted CV generation, interview preparation, and application tracking."
)

st.markdown(
    f"""
<div class="workflow">
  <div class="workflow-step"><span class="step-badge {"done" if step > 1 else "active"}">1</span><span>Analyze</span></div>
  <div class="workflow-line"></div>
  <div class="workflow-step"><span class="step-badge {"active" if step == 2 else ""}">2</span><span>Optimize & Export</span></div>
</div>
""",
    unsafe_allow_html=True,
)

with st.container(border=True):
    h1, h2, h3, h4 = st.columns([1.3, 1.25, 1.0, 0.85])
    h1.markdown(
        f'<div class="ui-card-title">AI engine</div><div class="ui-card-value">{cfg["label"]}</div>'
        f'<div class="ui-card-sub">{cfg["route_mode"].title()} route · {cfg["status"]}</div>',
        unsafe_allow_html=True,
    )
    h2.markdown(
        f'<div class="ui-card-title">Provider / model</div><div class="ui-card-value">{PROVIDER_LABELS.get(cfg["provider"], cfg["provider"])}</div>'
        f'<div class="ui-card-sub"><code>{cfg["model_name"]}</code></div>',
        unsafe_allow_html=True,
    )

    quota_snapshot = fetch_quota_status(st.session_state.get("api_base", DEFAULT_API_BASE)) or {}
    active_quota = quota_snapshot.get(cfg["provider"], {})
    active_minute = (
        (active_quota.get("windows", {}) or {}).get("minute", {}) if active_quota else {}
    )
    rpm_limit = (active_quota.get("limits", {}) or {}).get("rpm", 0) if active_quota else 0
    rpm_used = active_minute.get("used_requests", 0) if active_minute else 0
    rpm_pct = (
        float(active_minute.get("percent_used_requests", 0.0) or 0.0) if active_minute else 0.0
    )
    h3.markdown(
        f'<div class="ui-card-title">Requests / minute</div><div class="ui-card-value">{rpm_used}{" / " + str(rpm_limit) if rpm_limit else ""}</div>'
        f'<div class="ui-card-sub">{"%.0f%% used" % rpm_pct if rpm_limit else "No declared RPM limit"}</div>',
        unsafe_allow_html=True,
    )

    status_text = (
        "Within limit"
        if not rpm_limit or rpm_pct < 85
        else "Approaching limit"
        if rpm_pct < 100
        else "At limit"
    )
    status_class = (
        "health-ok"
        if not rpm_limit or rpm_pct < 85
        else "health-warn"
        if rpm_pct < 100
        else "health-bad"
    )
    h4.markdown(
        f'<div class="ui-card-title">Status</div><div class="{status_class}">{status_text}</div>'
        f'<div class="ui-card-sub">Backend connected at {ACTIVE_API_BASE}</div>',
        unsafe_allow_html=True,
    )

with st.expander("Usage, limits & provider health", expanded=False):
    usage_left, quota_right = st.columns(2)

    with usage_left:
        st.markdown("**Token usage**")
        period_labels = {
            "today": "Today",
            "7d": "Last 7 days",
            "30d": "Last 30 days",
            "all": "All time",
        }
        period_choice = st.selectbox(
            "Time window",
            list(period_labels.keys()),
            index=3,
            format_func=lambda x: period_labels[x],
            key="usage_period_select_main",
        )
        summary = fetch_usage_summary(
            st.session_state.get("api_base", DEFAULT_API_BASE), period_choice
        )
        if not summary:
            st.info("No usage data available yet.")
        else:
            u1, u2, u3 = st.columns(3)
            u1.metric("Tokens", fmt_tokens(summary.get("total_tokens", 0)))
            u2.metric("API calls", summary.get("success_calls", 0))
            u3.metric(
                "Est. cost",
                f"${summary.get('estimated_cost_usd', 0.0) or 0.0:.4f}",
            )
            by_model = summary.get("by_model") or {}
            if by_model:
                rows = [
                    {
                        "Model": model,
                        "Calls": stats.get("calls", 0),
                        "Tokens": fmt_tokens(stats.get("total_tokens", 0)),
                        "Cost": f"${stats.get('estimated_cost_usd', 0.0):.4f}",
                    }
                    for model, stats in sorted(
                        by_model.items(),
                        key=lambda kv: kv[1].get("total_tokens", 0),
                        reverse=True,
                    )
                ]
                st.dataframe(rows, width="stretch", hide_index=True)
        if st.button("Refresh usage", key="refresh_usage_main"):
            st.cache_data.clear()
            st.rerun()

    with quota_right:
        st.markdown("**Rate limits**")
        if not quota_snapshot:
            st.info("No quota data available. Run a call and refresh.")
        else:
            for prov in [p for p in quota_snapshot.keys() if p != "ollama"]:
                data = quota_snapshot.get(prov) or {}
                limits = data.get("limits", {}) or {}
                minute = (data.get("windows", {}) or {}).get("minute", {}) or {}
                day = (data.get("windows", {}) or {}).get("day", {}) or {}
                hits = data.get("rate_limit_hits_24h", 0)
                st.markdown(
                    f'<div class="quota-card"><div class="quota-title">{PROVIDER_LABELS.get(prov, prov.title())}</div>'
                    f'<div class="quota-meta"><code>{data.get("model") or "default"}</code> · '
                    f"{'429 hits: ' + str(hits) if hits else 'No 429 hits in 24h'}</div></div>",
                    unsafe_allow_html=True,
                )
                for label, window, key in [
                    ("Requests / minute", minute, "rpm"),
                    ("Requests / 24h", day, "rpd"),
                ]:
                    if limits.get(key, 0):
                        pct = float(window.get("percent_used_requests", 0.0) or 0.0)
                        used, lim, remaining = (
                            window.get("used_requests", 0),
                            window.get("limit_requests", 0),
                            window.get("remaining_requests", 0),
                        )
                        color = "#10b981" if pct < 60 else "#f59e0b" if pct < 85 else "#ef4444"
                        st.markdown(
                            f'<div class="quota-meta">{label} · {pct:.0f}% used</div>'
                            f'<div class="quota-bar-bg"><div class="quota-bar-fill" style="width:{max(0, min(100, pct)):.1f}%;background:{color};"></div></div>'
                            f'<div class="quota-meta">{used}/{lim} used · {remaining} remaining · resets {fmt_reset_local(window.get("resets_at", ""))}</div>',
                            unsafe_allow_html=True,
                        )
        st.caption(
            "Local estimate from this app's log and declared limits; not an account-wide provider balance."
        )
        if st.button("Refresh limits", key="refresh_quota_main"):
            st.cache_data.clear()
            st.rerun()

# ==============================================================================
# STEP 1 — INPUT
# ==============================================================================

with st.container(border=True):
    st.subheader("1. Analyze your resume against a target job")
    st.caption(
        "Provide the two inputs below. The resulting analysis powers the CV generator and career tools."
    )

    input_col, file_col = st.columns([1.45, 1], gap="large")
    with input_col:
        job_desc = st.text_area(
            "Target job description",
            height=235,
            placeholder="Paste the complete job description here…",
            value=st.session_state.get("job_desc", ""),
        )
        st.caption(
            "Include responsibilities, required skills, technologies, and qualifications where possible."
        )

    with file_col:
        uploaded_file = st.file_uploader(
            "Current resume",
            type=["pdf", "docx", "txt"],
            help="Supported formats: PDF, DOCX, TXT.",
        )
        if uploaded_file:
            size_kb = len(uploaded_file.getvalue()) / 1024
            st.success(f"Ready: {uploaded_file.name}")
            st.caption(f"{size_kb:.0f} KB · {uploaded_file.type or 'document'}")
        else:
            st.info("Upload one resume to begin.")

    ready = bool(uploaded_file and job_desc.strip())
    if not ready:
        missing = []
        if not uploaded_file:
            missing.append("resume")
        if not job_desc.strip():
            missing.append("job description")
        st.caption("Still needed: " + " and ".join(missing) + ".")

    if st.button(
        "Analyze ATS match & skill gaps",
        width="stretch",
        type="primary",
        disabled=not ready,
    ):
        file_bytes = uploaded_file.getvalue()
        st.session_state["uploaded_file_data"] = (
            uploaded_file.name,
            file_bytes,
            uploaded_file.type or "application/octet-stream",
        )
        st.session_state["job_desc"] = job_desc
        response = api_request(
            ENDPOINTS["analyze"],
            data=llm_request_data(job_description=job_desc),
            files={
                "resume_file": (
                    uploaded_file.name,
                    file_bytes,
                    uploaded_file.type or "application/octet-stream",
                )
            },
            spinner_text=f"Analyzing with {cfg['label']}…",
        )
        if response is None:
            st.error(f"Could not reach FastAPI backend at {ACTIVE_API_BASE}")
        elif response.status_code == 200:
            try:
                st.session_state["last_analysis"] = response.json()
                st.rerun()
            except Exception as exc:
                st.error(f"Backend returned invalid JSON: {exc}")
        else:
            st.error(f"Analysis failed: {error_detail(response)}")


# ==============================================================================
# BACKEND LOG (fragment)
# ==============================================================================


@st.fragment
def _render_log_panel():
    with st.expander("🛠️ Backend processing log", expanded=False):
        if st.button("Clear Logs", key="clear_logs_btn"):
            try:
                del_res = requests.delete(ENDPOINTS["processing_log"], timeout=10)
                if del_res.status_code == 200:
                    st.success("Backend log cleared.")
                    st.cache_data.clear()
                    st.rerun()
                else:
                    st.error(f"Clear failed: HTTP {del_res.status_code}")
            except requests.RequestException as exc:
                st.error(f"Clear failed: {exc}")

        logs = fetch_processing_log(
            st.session_state.get("api_base", DEFAULT_API_BASE),
            50,
        )

        log_file_path = "Log path not available"
        try:
            status_res = requests.get(
                f"{ACTIVE_API_BASE.rstrip('/')}/api/v1/resume/backend-status",
                timeout=5,
            )
            if status_res.status_code == 200:
                log_file_path = status_res.json().get("log_file") or log_file_path
        except Exception:
            pass

        if logs:
            all_kinds = sorted({(e.get("kind") or "llm") for e in logs})
            kind_filter = st.selectbox(
                "Filter by kind",
                ["All"] + all_kinds,
                index=0,
                key="log_kind_filter",
            )
            filtered = (
                logs
                if kind_filter == "All"
                else [e for e in logs if (e.get("kind") or "llm") == kind_filter]
            )
            rows = [
                {
                    "Time (UTC)": e.get("timestamp", ""),
                    "Kind": e.get("kind", "llm"),
                    "Event": e.get("event", ""),
                    "Detail": _fmt_log_detail(e),
                    "Status": e.get("status", ""),
                    "Duration ms": str(e.get("duration_ms") or ""),
                    "Error": str(e.get("error") or "")[:200],
                }
                for e in reversed(filtered)
            ]
            st.dataframe(
                pd.DataFrame(rows, dtype=str),
                width="stretch",
                hide_index=True,
            )
            st.caption(
                f"Backend log file: {log_file_path} · " f"{len(filtered)}/{len(logs)} events shown"
            )
        else:
            st.info(
                "No backend LLM processing events yet. "
                "Run an analysis first — log entries are written on every LLM call."
            )


_render_log_panel()

# ==============================================================================
# RESULTS
# ==============================================================================

if st.session_state.get("last_analysis"):
    result = st.session_state["last_analysis"]
    with st.container(border=True):
        st.subheader("2. Analysis results")
        m1, m2, m3 = st.columns(3)
        with m1:
            score_block("ATS Match Score", result.get("ats_match_score", 0))
        with m2:
            score_block("Keyword Density", result.get("keyword_density_score", 0))
        with m3:
            st.markdown(
                '<div class="metric-label">Missing Skills</div>',
                unsafe_allow_html=True,
            )
            st.markdown(
                f'<div class="metric-big">{len(result.get("missing_skills", []))}</div>',
                unsafe_allow_html=True,
            )
        st.divider()
        c1, c2 = st.columns(2)
        with c1:
            st.markdown("**Matching Skills**")
            chips(result.get("matching_skills", []), "good")
        with c2:
            st.markdown("**Missing Skills**")
            chips(result.get("missing_skills", []), "bad")

    with st.container(border=True):
        st.subheader("Priority improvements")
        for suggestion in result.get("improvement_suggestions", []) or [
            "No improvement suggestions were returned."
        ]:
            st.info(suggestion)

    recommendation = result.get("recommendation") or {}
    if recommendation:
        with st.container(border=True):
            st.subheader("Recommended CV format")
            label = recommendation.get("label", "Standard")
            reason = recommendation.get("reason", "")
            if recommendation.get("language_mismatch", False):
                st.warning(f"**{label}** — {reason}")
            else:
                st.success(f"**{label}** — {reason}")

    with st.expander("Export analysis report"):
        a, b = st.columns(2)
        with a:
            st.download_button(
                "Download TXT",
                generate_txt_export(result),
                "resume_analysis_report.txt",
                "text/plain",
                width="stretch",
            )
        with b:
            st.download_button(
                "Download DOCX",
                generate_docx_export(result),
                "resume_analysis_report.docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                width="stretch",
            )


# ==============================================================================
# CV GENERATION (fragment)
# ==============================================================================


@st.fragment
def _render_cv_generation():
    st.divider()
    st.header("Generate your tailored CV")

    _actionable_preview = actionable_suggestions_from_last_analysis()
    if _actionable_preview:
        with st.expander(
            f"📋 {len(_actionable_preview)} ATS improvements will be enforced in this CV",
            expanded=False,
        ):
            st.caption(
                "These items come from the analysis above and are passed as hard "
                "constraints into every CV generation call below."
            )
            for s in _actionable_preview:
                st.markdown(f"- {s}")
    else:
        st.caption("No actionable improvements were detected — universal ATS rules will apply.")

    tab1, tab2 = st.tabs(["📄 Standard ATS Resume", "🇩🇪 German Lebenslauf / PDF Options"])

    with tab1:
        st.caption(
            "Layout is auto-selected from the job description's language "
            "(English JD → compact single-page English layout)."
        )
        if st.button(
            "🪄 Build Tailored Resume",
            width="stretch",
            type="primary",
            key="build_docx_btn",
        ):
            if not build_files_payload():
                st.error("Resume file is no longer available. Please upload it again.")
            else:
                response = api_request(
                    ENDPOINTS["full_docx"],
                    data=llm_request_data(
                        job_description=st.session_state["job_desc"],
                        layout_style="auto",
                    ),
                    files=build_files_payload(),
                    spinner_text="Building tailored resume...",
                )
                if response and response.status_code == 200:
                    st.download_button(
                        "📥 Download Resume (.DOCX)",
                        response.content,
                        "Optimized_Tailored_Resume.docx",
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        width="stretch",
                    )
                else:
                    st.error(f"Generation failed: {error_detail(response)}")

    with tab2:
        template_options = list(TEMPLATE_LABELS)
        selected_layout = st.selectbox(
            "Layout Style",
            template_options,
            index=0,
            format_func=lambda x: TEMPLATE_LABELS[x],
            help="Auto picks the layout that matches the job description's language.",
            key="cv_layout_select",
        )
        common_data = llm_request_data(
            job_description=st.session_state["job_desc"],
            layout_style=selected_layout,
            template_style=selected_layout,
        )
        b1, b2 = st.columns(2)
        with b1:
            if st.button(
                "📄 Build PDF",
                width="stretch",
                type="primary",
                key="build_pdf_btn",
            ):
                response = api_request(
                    ENDPOINTS["german_pdf"],
                    common_data,
                    build_files_payload(),
                    "Generating PDF...",
                )
                if response and response.status_code == 200:
                    st.download_button(
                        "📥 Download PDF",
                        response.content,
                        f"CV_{selected_layout}.pdf",
                        "application/pdf",
                        width="stretch",
                    )
                else:
                    st.error(f"PDF generation failed: {error_detail(response)}")
        with b2:
            if st.button("🛠️ Build TEX Source", width="stretch", key="build_tex_btn"):
                response = api_request(
                    ENDPOINTS["german_tex"],
                    common_data,
                    build_files_payload(),
                    "Generating LaTeX source...",
                )
                if response and response.status_code == 200:
                    st.download_button(
                        "📥 Download TEX",
                        response.content,
                        f"CV_{selected_layout}.tex",
                        "text/plain",
                        width="stretch",
                    )
                else:
                    st.error(f"TEX generation failed: {error_detail(response)}")


if st.session_state.get("last_analysis"):
    _render_cv_generation()


# ==============================================================================
# ADVANCED TOOLS (fragment)
# ==============================================================================


@st.fragment
def _render_advanced_tools():
    st.divider()
    st.header("Advanced tools")
    tool_tab1, tool_tab2 = st.tabs(["🔍 Visual Bullet Diff", "👥 Bulk CV Screening"])

    with tool_tab1:
        orig_text = st.text_area("Original Bullets (one per line)", height=100, key="diff_orig")
        opt_text = st.text_area("Optimized Bullets (one per line)", height=100, key="diff_opt")
        if st.button("Compare Bullets", width="stretch", key="diff_compare_btn"):
            payload = {
                "original_bullets": [x.strip() for x in orig_text.splitlines() if x.strip()],
                "optimized_bullets": [x.strip() for x in opt_text.splitlines() if x.strip()],
            }
            try:
                response = requests.post(
                    ENDPOINTS["diff_preview"],
                    json=payload,
                    timeout=REQUEST_TIMEOUT,
                )
                if response.status_code == 200:
                    diffs = response.json().get("diffs", [])
                    render_diff_view(diffs)
                else:
                    st.error(error_detail(response))
            except requests.RequestException as exc:
                st.error(str(exc))

    with tool_tab2:
        bulk_job_desc = st.text_area(
            "Target Job Description for Bulk Screening",
            height=120,
            key="bulk_jd",
        )
        bulk_files = st.file_uploader(
            "Upload Multiple Resumes",
            type=["pdf", "docx", "txt"],
            accept_multiple_files=True,
            key="bulk_files_uploader",
        )
        if st.button(
            "Run Bulk Analysis",
            type="primary",
            width="stretch",
            key="bulk_run_btn",
        ):
            if not bulk_job_desc.strip() or not bulk_files:
                st.warning("Please provide both a job description and at least one resume file.")
            else:
                files_payload = [
                    (
                        "resume_files",
                        (
                            f.name,
                            f.getvalue(),
                            f.type or "application/octet-stream",
                        ),
                    )
                    for f in bulk_files
                ]
                response = api_request(
                    ENDPOINTS["analyze_bulk"],
                    llm_request_data(job_description=bulk_job_desc),
                    files_payload,
                    "Screening applicant batch...",
                )
                if response and response.status_code == 200:
                    results = response.json().get("rankings", [])
                    st.success(f"Successfully processed {len(results)} candidate resumes.")
                    df = pd.DataFrame(results)
                    if not df.empty:
                        cols = [
                            c
                            for c in ["filename", "ats_score", "keyword_density"]
                            if c in df.columns
                        ]
                        st.dataframe(df[cols], width="stretch", hide_index=True)
                else:
                    st.error(f"Bulk screening failed: {error_detail(response)}")


_render_advanced_tools()


# ==============================================================================
# CAREER SUITE (fragment)
# ==============================================================================


@st.fragment
def _render_career_suite():
    st.divider()
    st.header("Career suite")
    f_tab1, f_tab2, f_tab3, f_tab4, f_tab5 = st.tabs(
        [
            "📊 Audit Matrix",
            "✉️ Cover Letter",
            "🎯 Interview Prep",
            "💼 LinkedIn Optimizer",
            "📌 Application Pipeline",
        ]
    )

    with f_tab1:
        if st.button("Run Detailed ATS Audit", key="run_audit_btn"):
            if not build_files_payload():
                st.warning("Please upload a resume in Step 1.")
            else:
                res = api_request(
                    ENDPOINTS["audit_matrix"],
                    llm_request_data(job_description=st.session_state.get("job_desc", "")),
                    build_files_payload(),
                    "Generating ATS Audit Matrix...",
                )
                if res and res.status_code == 200:
                    st.json(res.json().get("data", {}))

    with f_tab2:
        comp_name = st.text_input("Company Name", value="Target Company", key="cover_company")
        cov_tone = st.selectbox("Tone", ["formal", "startup", "technical"], key="cover_tone")
        if st.button("Generate Outreach Materials", key="gen_cover_btn"):
            res = api_request(
                ENDPOINTS["cover_letter"],
                llm_request_data(
                    job_description=st.session_state.get("job_desc", ""),
                    company_name=comp_name,
                    tone=cov_tone,
                ),
                build_files_payload(),
                "Generating Cover Letter & Cold Email...",
            )
            if res and res.status_code == 200:
                data = res.json().get("data", {})
                st.subheader("Cover Letter")
                st.info(data.get("cover_letter"))
                st.subheader("Cold Outreach Message")
                st.success(data.get("cold_outreach"))

    with f_tab3:
        if st.button("Generate Interview Strategy", key="gen_interview_btn"):
            res = api_request(
                ENDPOINTS["interview_prep"],
                llm_request_data(job_description=st.session_state.get("job_desc", "")),
                build_files_payload(),
                "Generating Interview Questions & Gap Defenses...",
            )
            if res and res.status_code == 200:
                st.json(res.json().get("data", {}))

    with f_tab4:
        target_role_input = st.text_input(
            "Target Role Title",
            value="Software Engineer",
            key="linkedin_role",
        )
        if st.button("Optimize LinkedIn Profile", key="opt_linkedin_btn"):
            res = api_request(
                ENDPOINTS["linkedin_optimize"],
                llm_request_data(target_role=target_role_input),
                build_files_payload(),
                "Generating LinkedIn Headlines & About Section...",
            )
            if res and res.status_code == 200:
                st.json(res.json().get("data", {}))

    with f_tab5:
        try:
            t_res = requests.get(ENDPOINTS["tracker"], timeout=10)
            if t_res.status_code == 200:
                apps = t_res.json().get("applications", [])
                if apps:
                    st.dataframe(
                        pd.DataFrame(apps),
                        width="stretch",
                        hide_index=True,
                    )
                else:
                    st.info("No applications saved yet in local SQLite tracker.")
        except Exception:
            st.caption("Application tracker pipeline ready.")


_render_career_suite()


# ==============================================================================
# ANALYTICS (fragment)
# ==============================================================================


@st.fragment
def _render_analytics():
    st.divider()
    st.header("Analytics")

    @st.cache_data(ttl=30)
    def _fetch_analytics(api_base: str, period: str) -> dict | None:
        try:
            url = f"{api_base.rstrip('/')}/api/v1/resume/analytics/summary"
            r = requests.get(url, params={"period": period}, timeout=15)
            if r.status_code == 200:
                return r.json().get("data") or {}
        except Exception:
            return None
        return None

    period_labels = {
        "today": "Today",
        "7d": "Last 7 days",
        "30d": "Last 30 days",
        "90d": "Last 90 days",
        "all": "All time",
    }
    period_choice = st.selectbox(
        "Time window",
        list(period_labels.keys()),
        index=2,
        format_func=lambda x: period_labels[x],
        key="analytics_period",
    )

    data = _fetch_analytics(
        st.session_state.get("api_base", DEFAULT_API_BASE),
        period_choice,
    )

    if not data:
        st.info("No analytics available for this window.")
        if st.button("Refresh", key="analytics_refresh_empty"):
            st.cache_data.clear()
            st.rerun()
        return

    meta = data.get("meta") or {}
    llm = data.get("llm") or {}
    apps = data.get("applications") or {}
    pipeline = data.get("pipeline") or {}
    skills = data.get("skills") or {}

    m1, m2, m3, m4 = st.columns(4)
    m1.metric("LLM calls", llm.get("total_calls", 0))
    m2.metric("Tokens", fmt_tokens(llm.get("total_tokens", 0)))
    m3.metric("Est. cost", f"${llm.get('total_cost_usd', 0.0):.4f}")
    m4.metric("Failure rate", f"{llm.get('failure_rate', 0.0) * 100:.1f}%")

    st.divider()

    left, right = st.columns(2)
    with left:
        st.markdown("**Token usage over time**")
        series = llm.get("tokens_by_day") or []
        if series:
            df = pd.DataFrame(series).set_index("date")
            st.bar_chart(df, height=220, color="#2563eb")
        else:
            st.caption("No token activity in this window.")

    with right:
        st.markdown("**Cost over time**")
        series = llm.get("cost_by_day") or []
        if series:
            df = pd.DataFrame(series).set_index("date")
            st.line_chart(df, height=220, color="#16a34a")
        else:
            st.caption("No cost data in this window.")

    st.divider()

    left, right = st.columns(2)
    with left:
        st.markdown("**Applications by status**")
        status_counts = apps.get("by_status") or {}
        if status_counts and any(status_counts.values()):
            df = pd.DataFrame(
                [{"status": k, "count": v} for k, v in status_counts.items()]
            ).set_index("status")
            st.bar_chart(df, height=220, color="#6366f1")
        else:
            st.caption("No applications recorded yet.")

    with right:
        st.markdown("**ATS score distribution**")
        buckets = apps.get("by_score_bucket") or {}
        if buckets and any(buckets.values()):
            df = pd.DataFrame(
                [{"score range": k, "count": v} for k, v in buckets.items()]
            ).set_index("score range")
            st.bar_chart(df, height=220, color="#f59e0b")
        else:
            st.caption("No ATS scores recorded yet.")

    st.divider()

    left, right = st.columns(2)
    with left:
        st.markdown("**Top missing skills**")
        top_missing = skills.get("top_missing") or []
        if top_missing:
            df = pd.DataFrame([{"skill": k, "count": v} for k, v in top_missing]).set_index("skill")
            st.bar_chart(df, height=280, color="#dc2626")
        else:
            st.caption("No missing-skill data yet.")

    with right:
        st.markdown("**LLM calls by provider**")
        providers = llm.get("calls_by_provider") or {}
        if providers:
            df = pd.DataFrame(
                [{"provider": k, "calls": v} for k, v in providers.items()]
            ).set_index("provider")
            st.bar_chart(df, height=280, color="#0891b2")
        else:
            st.caption("No provider data yet.")

    st.divider()

    left, right = st.columns([1, 1.4])
    with left:
        st.markdown("**Avg operation duration (ms)**")
        durations = pipeline.get("avg_duration_ms_by_op") or {}
        if durations:
            rows = [
                {"operation": op, "avg_ms": v}
                for op, v in sorted(durations.items(), key=lambda kv: -kv[1])
            ]
            st.dataframe(rows, width="stretch", hide_index=True)
        else:
            st.caption("No pipeline duration data yet.")

    with right:
        st.markdown("**Recent errors**")
        errors = data.get("recent_errors") or []
        if errors:
            rows = [
                {
                    "Time": e.get("timestamp", "")[:19].replace("T", " "),
                    "Kind": e.get("kind", ""),
                    "Provider": e.get("provider", ""),
                    "Error": (e.get("error") or "")[:120],
                }
                for e in errors
            ]
            st.dataframe(
                pd.DataFrame(rows, dtype=str),
                width="stretch",
                hide_index=True,
            )
        else:
            st.caption("No errors in this window.")

    st.caption(
        f"Window: {period_labels[period_choice]} · "
        f"{meta.get('events_count', 0)} events · "
        f"{meta.get('applications_count', 0)} applications"
    )

    if st.button("Refresh analytics", key="analytics_refresh"):
        st.cache_data.clear()
        st.rerun()


_render_analytics()
