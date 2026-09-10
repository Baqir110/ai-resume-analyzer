import io
from pypdf import PdfReader
from docx import Document
from fastapi import UploadFile, HTTPException


async def extract_text_from_file(file: UploadFile) -> str:
    import time as _time
    from app.core.event_log import log_event

    _started = _time.perf_counter()
    _fname = file.filename or ""
    _ext = _fname.rsplit(".", 1)[-1].lower() if "." in _fname else "unknown"

    try:
        text = await _extract_text_impl(file)
    except HTTPException as exc:
        log_event(
            "parse",
            "parse_failed",
            file_type=_ext,
            filename=_fname,
            duration_ms=round((_time.perf_counter() - _started) * 1000, 1),
            error=str(getattr(exc, "detail", exc))[:300],
        )
        raise
    except Exception as exc:
        log_event(
            "parse",
            "parse_failed",
            file_type=_ext,
            filename=_fname,
            duration_ms=round((_time.perf_counter() - _started) * 1000, 1),
            error=str(exc)[:300],
        )
        raise

    log_event(
        "parse",
        "parse_completed",
        file_type=_ext,
        filename=_fname,
        chars_extracted=len(text or ""),
        empty=not bool((text or "").strip()),
        duration_ms=round((_time.perf_counter() - _started) * 1000, 1),
    )
    return text


async def _extract_text_impl(file: UploadFile) -> str:
    try:
        # Read file bytes
        content = await file.read()
        # Reset file pointer so other downstream functions can re-read if needed
        await file.seek(0)
    except Exception as e:
        raise HTTPException(
            status_code=400, detail=f"Failed to read upload stream: {str(e)}"
        )

    filename = (file.filename or "").lower()

    # 1. PDF Files
    if filename.endswith(".pdf"):
        try:
            reader = PdfReader(io.BytesIO(content))
            pages_text = []
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    pages_text.append(extracted.strip())
            return "\n\n".join(pages_text).strip()
        except Exception as e:
            raise HTTPException(
                status_code=400, detail=f"Failed to parse PDF file: {str(e)}"
            )

    # 2. DOCX Files (Paragraphs + Tables)
    elif filename.endswith(".docx"):
        try:
            doc = Document(io.BytesIO(content))
            extracted_blocks = []

            # Extract standard paragraphs
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    extracted_blocks.append(text)

            # Extract content from tables if present
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        cell_text = cell.text.strip()
                        if cell_text and cell_text not in extracted_blocks:
                            extracted_blocks.append(cell_text)

            return "\n\n".join(extracted_blocks).strip()
        except Exception as e:
            raise HTTPException(
                status_code=400, detail=f"Failed to parse DOCX file: {str(e)}"
            )

    # 3. Plain Text Files
    elif filename.endswith(".txt"):
        try:
            return content.decode("utf-8", errors="ignore").strip()
        except Exception as e:
            raise HTTPException(
                status_code=400, detail=f"Failed to parse TXT file: {str(e)}"
            )

    # 4. Fallback for Unsupported Types
    else:
        raise HTTPException(
            status_code=400,
            detail="Unsupported file format. Please upload a PDF, DOCX, or TXT file.",
        )
